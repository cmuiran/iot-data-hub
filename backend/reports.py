"""Period reports over several tags: hourly / daily / monthly, in the plant's time zone.

Functions per tag:

* ``avg`` — mean of the stored good values (weighted by sample count, not by time)
* ``min`` / ``max`` — extremes
* ``last`` — last value of the period
* ``delta`` — consumption of a counter (energy, run hours, volume): the increase over the
  period, built from hour to hour. A step down (counter reset or meter exchange) counts only the
  increase inside that hour, so a reset never produces a negative consumption.

Days and months start at local midnight (``HUB_TIMEZONE``, default Asia/Tehran), so a daily
energy report matches the plant's calendar, not UTC.
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from datetime import datetime, timedelta
from typing import Any, Literal
from zoneinfo import ZoneInfo

from pydantic import BaseModel, Field, field_validator

Interval = Literal["hour", "day", "month"]
Function = Literal["avg", "min", "max", "last", "delta"]
MAX_TAGS = 50
MAX_ROWS = 10_000


class ReportColumn(BaseModel):
    tag: str = Field(min_length=7, max_length=263)
    function: Function = "avg"
    label: str | None = Field(default=None, max_length=100)


class ReportRequest(BaseModel):
    columns: list[ReportColumn] = Field(min_length=1, max_length=MAX_TAGS)
    start: datetime
    end: datetime
    interval: Interval = "day"
    title: str | None = Field(default=None, max_length=200)

    @field_validator("start", "end")
    @classmethod
    def _aware(cls, value: datetime) -> datetime:
        if value.tzinfo is None:
            raise ValueError("needs a UTC offset (e.g. 2026-10-01T00:00:00+03:30)")
        return value


@dataclass
class _Hour:
    bucket: datetime
    avg: float | None
    min: float | None
    max: float | None
    first: float | None
    last: float | None
    n: int


def period_start(ts: datetime, interval: Interval, tz: ZoneInfo) -> datetime:
    local = ts.astimezone(tz)
    if interval == "hour":
        local = local.replace(minute=0, second=0, microsecond=0)
    elif interval == "day":
        local = local.replace(hour=0, minute=0, second=0, microsecond=0)
    else:
        local = local.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    return local


def periods(start: datetime, end: datetime, interval: Interval, tz: ZoneInfo) -> list[datetime]:
    """Local period starts covering [start, end)."""
    out: list[datetime] = []
    current = period_start(start, interval, tz)
    while current < end:
        out.append(current)
        if interval == "hour":
            current = (current.astimezone(ZoneInfo("UTC")) + timedelta(hours=1)).astimezone(tz)
        elif interval == "day":
            current = current + timedelta(days=1)  # wall-clock arithmetic: next local midnight
        else:
            year, month = (current.year + 1, 1) if current.month == 12 else (current.year,
                                                                               current.month + 1)
            current = current.replace(year=year, month=month)
        if len(out) > MAX_ROWS:
            raise ValueError(f"more than {MAX_ROWS} rows; choose a longer interval")
    return out


def aggregate(hours: list[_Hour], function: Function, previous_last: float | None
              ) -> tuple[float | None, float | None]:
    """Value of one period, and the last value to carry into the next period (for delta)."""
    good = [h for h in hours if h.n]
    if not good:
        return None, previous_last
    last = good[-1].last
    if function == "avg":
        total = sum((h.avg or 0) * h.n for h in good)
        return total / sum(h.n for h in good), last
    if function == "min":
        return min(h.min for h in good if h.min is not None), last
    if function == "max":
        return max(h.max for h in good if h.max is not None), last
    if function == "last":
        return last, last
    consumed, prev = 0.0, previous_last
    for h in good:
        if h.first is None or h.last is None:
            continue
        step = h.last - prev if prev is not None else h.last - h.first
        if step < 0:  # counter reset inside or before this hour
            step = max(h.last - h.first, 0.0)
        consumed += step
        prev = h.last
    return consumed, prev


async def build_report(store: Any, request: ReportRequest, tz_name: str,
                       units: dict[str, str | None]) -> dict[str, Any]:
    tz = ZoneInfo(tz_name)
    if request.start >= request.end:
        raise ValueError("start must be before end")
    starts = periods(request.start, request.end, request.interval, tz)
    bounds = [(s, starts[i + 1] if i + 1 < len(starts) else request.end)
              for i, s in enumerate(starts)]
    columns = []
    for col in request.columns:
        # one hour of look-back so the first period's delta starts from the previous reading
        raw = await store.hourly_series(col.tag, request.start - timedelta(hours=1), request.end,
                                        tz_name)
        hours = [_Hour(r["bucket"], r["avg_value"], r["min_value"], r["max_value"],
                       r["first_value"], r["last_value"], int(r["n"])) for r in raw]
        before = [h for h in hours if h.bucket < request.start]
        carry = before[-1].last if before and col.function == "delta" else None
        values = []
        for lo, hi in bounds:
            inside = [h for h in hours if lo <= h.bucket < hi and h.bucket >= request.start]
            value, carry = aggregate(inside, col.function, carry)
            values.append(value)
        columns.append({"tag": col.tag, "function": col.function,
                        "label": col.label or f"{col.tag.rsplit('.', 1)[1]} ({col.function})",
                        "unit": units.get(col.tag), "values": values,
                        "total": (sum(v for v in values if v is not None)
                                  if col.function == "delta" else None)})
    return {"title": request.title, "interval": request.interval, "timezone": tz_name,
            "start": request.start.isoformat(), "end": request.end.isoformat(),
            "periods": [s.isoformat() for s in starts], "columns": columns}


def to_csv(report: dict[str, Any]) -> str:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["period"] + [f"{c['label']} [{c['unit'] or ''}]" for c in report["columns"]])
    for i, period in enumerate(report["periods"]):
        writer.writerow([period] + ["" if c["values"][i] is None else round(c["values"][i], 6)
                                    for c in report["columns"]])
    if any(c["total"] is not None for c in report["columns"]):
        writer.writerow(["total"] + ["" if c["total"] is None else round(c["total"], 6)
                                     for c in report["columns"]])
    return buf.getvalue()


def to_xlsx(report: dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1F4E79")
    row = 1
    if report.get("title"):
        ws.cell(row=1, column=1, value=report["title"]).font = Font(bold=True, size=14)
        row = 2
    ws.cell(row=row, column=1, value=f"{report['start']} - {report['end']} "
                                     f"({report['interval']}, {report['timezone']})")
    row += 2
    headers = ["Period"] + [c["label"] for c in report["columns"]]
    units = [""] + [c["unit"] or "" for c in report["columns"]]
    for col, (h, u) in enumerate(zip(headers, units, strict=True), start=1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.font, cell.fill = head, fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.cell(row=row + 1, column=col, value=u).alignment = Alignment(horizontal="center")
        ws.column_dimensions[cell.column_letter].width = 22 if col == 1 else 18
    row += 2
    fmt = {"hour": "%Y-%m-%d %H:%M", "day": "%Y-%m-%d", "month": "%Y-%m"}[report["interval"]]
    for i, period in enumerate(report["periods"]):
        ws.cell(row=row, column=1, value=datetime.fromisoformat(period).strftime(fmt))
        for col, c in enumerate(report["columns"], start=2):
            cell = ws.cell(row=row, column=col, value=c["values"][i])
            cell.number_format = "0.000"
        row += 1
    if any(c["total"] is not None for c in report["columns"]):
        ws.cell(row=row, column=1, value="Total").font = Font(bold=True)
        for col, c in enumerate(report["columns"], start=2):
            if c["total"] is not None:
                cell = ws.cell(row=row, column=col, value=c["total"])
                cell.font, cell.number_format = Font(bold=True), "0.000"
    ws.freeze_panes = ws.cell(row=5 if report.get("title") else 4, column=2)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
