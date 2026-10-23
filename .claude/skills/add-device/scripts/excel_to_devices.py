#!/usr/bin/env python3
"""Convert a devices workbook into config/devices/<device_id>.yaml files.

Sheets (see templates/devices_template.xlsx): Devices, ModbusRegisters, MqttTags, OpcuaTags,
S7Tags, HttpTags, SnmpTags. A ``states`` cell labels values: ``0=stopped; 1=running``.
Rows whose first cell starts with '#' are comments and are skipped. Empty cells are omitted so
the model defaults apply. A cell containing TODO means "unknown": the tag is left out (and
listed in the file header) or, for host / unit_id / endpoint_url only, a visible placeholder is
written and the device is forced to ``enabled: false``. Nothing is ever guessed.

Run from the project root with the project's Python (>= 3.12; needs openpyxl, a dev-only
helper; PyYAML is already a dependency). <skill-dir> is wherever the add-device skill is installed:

    python <skill-dir>/scripts/excel_to_devices.py devices.xlsx [--out config/devices]
        [--only ID ...] [--force] [--dry-run] [--check]

Before anything is written, every generated device is validated together with the existing
config (``hub.config_loader``, imported from the current directory). A device that would add a
validation ERROR — duplicate device_id, Modbus host:port:unit_id clash, overlapping MQTT topic,
unknown broker_ref, schema error — is not written, so one bad row can never stop the hub from
starting. ``--check`` additionally runs ``python -m hub.validate_config`` afterwards.

Exit status: 0 ok, 1 errors (some devices/tags not generated or validation failed), 2 usage.
"""

from __future__ import annotations

import argparse
import os
import re
import shutil
import subprocess
import sys
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import yaml

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is required for this helper: pip install openpyxl")

PROTOCOLS = ("modbus_tcp", "modbus_rtu", "mqtt", "opcua", "s7", "http", "snmp")
TAG_SHEETS = ("ModbusRegisters", "MqttTags", "OpcuaTags", "S7Tags", "HttpTags", "SnmpTags")
TODO_RE = re.compile(r"^\s*todo\b", re.IGNORECASE)
_DIGITS = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")

# Placeholders allowed for unknown connection values (device is then disabled).
PLACEHOLDER_HOST = "todo.invalid"
PLACEHOLDER_UNIT_ID = 0
PLACEHOLDER_ENDPOINT = "opc.tcp://todo.invalid:4840"

MODICON_AREAS = {"0": "coil", "1": "discrete_input", "3": "input", "4": "holding"}


class Unknown:
    """Marker for a cell that holds TODO."""

    def __repr__(self) -> str:
        return "TODO"


TODO = Unknown()


@dataclass
class Result:
    device_id: str
    data: dict[str, Any] | None = None
    todos: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)


# ---------------------------------------------------------------------------------------------
# Reading cells
# ---------------------------------------------------------------------------------------------


def clean(value: Any) -> Any:
    """Normalise one cell: None for blank, TODO marker, ASCII digits, trimmed strings."""
    if value is None:
        return None
    if isinstance(value, str):
        text = value.translate(_DIGITS).strip()
        if not text:
            return None
        if TODO_RE.match(text):
            return TODO
        return text
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def read_sheet(wb: Any, name: str) -> list[tuple[int, dict[str, Any]]]:
    """Rows as (excel row number, {column: cleaned value}); missing sheet -> []."""
    if name not in wb.sheetnames:
        return []
    rows = wb[name].iter_rows(values_only=True)
    header = next(rows, None)
    if header is None:
        return []
    columns = [str(h).strip().lower() if h is not None else "" for h in header]
    out: list[tuple[int, dict[str, Any]]] = []
    for offset, row in enumerate(rows, start=2):
        first = row[0] if row else None
        if first is None or (isinstance(first, str) and first.strip().startswith("#")):
            continue
        record = {c: clean(v) for c, v in zip(columns, row, strict=False) if c}
        if any(v is not None for v in record.values()):
            out.append((offset, record))
    return out


def as_bool(value: Any, where: str, res: Result) -> bool | None:
    if value is None or value is TODO:
        return None
    if isinstance(value, bool):
        return value
    text = str(value).strip().lower()
    if text in {"1", "true", "yes", "y"}:
        return True
    if text in {"0", "false", "no", "n"}:
        return False
    res.errors.append(f"{where}: {value!r} is not a boolean")
    return None


def as_number(value: Any, where: str, res: Result, integer: bool = False) -> Any:
    if value is None or value is TODO:
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        res.errors.append(f"{where}: {value!r} is not a number")
        return None
    if integer:
        if not number.is_integer():
            res.errors.append(f"{where}: {value!r} must be an integer")
            return None
        return int(number)
    return int(number) if number.is_integer() else number


# ---------------------------------------------------------------------------------------------
# Modbus address conversion
# ---------------------------------------------------------------------------------------------


def parse_int(text: Any) -> int | None:
    if isinstance(text, int) and not isinstance(text, bool):
        return text
    if isinstance(text, float) and text.is_integer():
        return int(text)
    if isinstance(text, str):
        try:
            return int(text, 16) if text.lower().startswith("0x") else int(text, 10)
        except ValueError:
            return None
    return None


def convert_modbus_address(
    raw: Any, style: str, register_type: str | None
) -> tuple[str | None, int | None, str | None]:
    """Return (register_type, 0-based address, error)."""
    number = parse_int(raw)
    if number is None:
        return register_type, None, f"address {raw!r} is not an integer"
    if style == "protocol":
        address, rtype = number, register_type
    elif style == "one_based":
        address, rtype = number - 1, register_type
    elif style == "modicon":
        # Keep the digits as typed: parse_int("00001") == 1 would lose the coil prefix '0'.
        digits = raw.strip() if isinstance(raw, str) else str(number)
        if not digits.isdigit() or len(digits) not in (5, 6):
            hint = ""
            if not isinstance(raw, str) and len(digits) < 5:
                hint = (
                    " - if this is a coil (00001-style), Excel dropped the leading zeros: "
                    "enter it as text ('00001) or use register_type coil + address_style one_based"
                )
            return register_type, None, (
                f"modicon address {raw!r} must have 5 or 6 digits (e.g. 40001, 00017){hint}; "
                "for short numbers use address_style one_based or protocol"
            )
        rtype = MODICON_AREAS.get(digits[0])
        if rtype is None:
            return register_type, None, f"modicon prefix {digits[0]!r} of {raw!r} is not 0/1/3/4"
        if register_type and register_type != rtype:
            return register_type, None, (
                f"register_type {register_type} contradicts modicon prefix of {raw!r} ({rtype})"
            )
        address = int(digits[1:]) - 1
    else:
        return register_type, None, f"address_style {style!r} must be protocol, one_based or modicon"
    if address < 0:
        return rtype, None, f"address {raw!r} converts to {address} (style {style}); check the base"
    if rtype is None:
        return None, None, "register_type is missing (needed unless address_style is modicon)"
    return rtype, address, None


# ---------------------------------------------------------------------------------------------
# Building device dictionaries
# ---------------------------------------------------------------------------------------------


def parse_states(value: Any, where: str, res: Result) -> dict[int, str] | None:
    """``0=stopped; 1=running`` (also ``,`` or new lines between items) -> {0: "stopped", ...}."""
    if value is None or value is TODO:
        return None
    out: dict[int, str] = {}
    for item in re.split(r"[;,\n]+", str(value)):
        if not item.strip():
            continue
        key, sep, label = item.partition("=")
        number = parse_int(key.strip().translate(_DIGITS))
        if not sep or number is None or not label.strip():
            res.errors.append(f"{where}: {item.strip()!r} is not value=label")
            return None
        out[number] = label.strip()
    return out or None


def put(target: dict[str, Any], key: str, value: Any) -> None:
    if value is not None and value is not TODO:
        target[key] = value


def build_tag(
    row: int, rec: dict[str, Any], sheet: str, res: Result, extras: dict[str, Any]
) -> dict[str, Any] | None:
    """Common tag fields. ``extras`` already holds the protocol-specific keys."""
    label = f"{sheet} row {row}"
    name = rec.get("tag")
    if name is None or name is TODO:
        res.todos.append(f"{label}: tag name is missing; row skipped")
        return None
    name = str(name).lower()
    where = f"{label} ({name})"
    if rec.get("data_type") in (None, TODO):
        res.todos.append(f"{where}: data_type is unknown; tag skipped")
        return None
    tag: dict[str, Any] = {"name": name}
    put(tag, "description", rec.get("description"))
    put(tag, "unit", rec.get("unit"))
    tag["data_type"] = str(rec["data_type"]).lower()
    tag.update(extras)
    scale = as_number(rec.get("scale"), f"{where}.scale", res)
    if scale is not None and scale != 1:
        tag["scale"] = scale
    offset = as_number(rec.get("offset"), f"{where}.offset", res)
    if offset is not None and offset != 0:
        tag["offset"] = offset
    put(tag, "min", as_number(rec.get("min"), f"{where}.min", res))
    put(tag, "max", as_number(rec.get("max"), f"{where}.max", res))
    history = as_bool(rec.get("history"), f"{where}.history", res)
    if history is False:
        tag["history"] = False
    put(tag, "history_deadband", as_number(rec.get("history_deadband"),
                                           f"{where}.history_deadband", res))
    states = parse_states(rec.get("states"), f"{where}.states", res)
    if states:
        tag["states"] = states
    access = rec.get("access")
    if access not in (None, TODO) and str(access).lower() != "read":
        tag["access"] = str(access).lower()
        if tag["access"] == "read_write":
            check = ("bool: on/off is the only range check" if tag["data_type"] == "bool"
                     else "confirm min/max with the datasheet")
            res.todos.append(
                f"{where}: WRITABLE tag - {check}; writes still need the control-safety "
                "layer (RBAC, audit)"
            )
    return tag


def modbus_tags(rows: list[tuple[int, dict[str, Any]]], res: Result) -> list[dict[str, Any]]:
    tags = []
    for row, rec in rows:
        raw = rec.get("address")
        label = f"ModbusRegisters row {row} ({rec.get('tag')})"
        if raw in (None, TODO):
            res.todos.append(f"{label}: address unknown; tag skipped")
            continue
        style = str(rec.get("address_style") or "protocol").lower()
        rtype = rec.get("register_type")
        rtype = None if rtype in (None, TODO) else str(rtype).lower()
        rtype, address, err = convert_modbus_address(raw, style, rtype)
        if err:
            res.errors.append(f"{label}: {err}")
            continue
        extras: dict[str, Any] = {"register_type": rtype, "address": address}
        bit = as_number(rec.get("bit"), f"{label}.bit", res, integer=True)
        if bit is not None:
            extras["bit"] = bit
        count = as_number(rec.get("count"), f"{label}.count", res, integer=True)
        if count is not None:
            extras["count"] = count
        bit_count = as_number(rec.get("bit_count"), f"{label}.bit_count", res, integer=True)
        if bit_count is not None:
            extras["bit_count"] = bit_count
        encoding = rec.get("encoding")
        if encoding not in (None, TODO) and str(encoding).lower() != "binary":
            extras["encoding"] = str(encoding).lower()
        tag = build_tag(row, rec, "ModbusRegisters", res, extras)
        if tag:
            tags.append(tag)
    return tags


def simple_tags(
    rows: list[tuple[int, dict[str, Any]]], sheet: str, key: str, res: Result,
    optional: tuple[str, ...] = (),
) -> list[dict[str, Any]]:
    """Tags addressed by one ``key`` column (json_path, node_id, address, oid); ``optional``
    columns are copied when filled (path, string_length)."""
    tags = []
    for row, rec in rows:
        value = rec.get(key)
        if value in (None, TODO):
            res.todos.append(f"{sheet} row {row} ({rec.get('tag')}): {key} unknown; tag skipped")
            continue
        extras: dict[str, Any] = {key: str(value)}
        for extra in optional:
            cell = rec.get(extra)
            if cell not in (None, TODO):
                extras[extra] = (as_number(cell, f"{sheet} row {row}.{extra}", res, integer=True)
                                 if extra == "string_length" else str(cell))
        tag = build_tag(row, rec, sheet, res, extras)
        if tag:
            tags.append(tag)
    return tags


def required_text(rec: dict[str, Any], key: str, where: str, res: Result) -> Any:
    value = rec.get(key)
    if value in (None, TODO):
        res.errors.append(f"{where}: {key} is required (got {'TODO' if value is TODO else 'blank'})")
        return None
    return str(value)


def build_device(
    row: int, rec: dict[str, Any], tags_by_device: dict[str, dict[str, list]]
) -> Result:
    device_id = str(rec.get("device_id") or "").lower()
    res = Result(device_id=device_id or f"(row {row})")
    where = f"Devices row {row} ({device_id})"
    if not device_id:
        res.errors.append(f"{where}: device_id is required")
        return res

    protocol = str(rec.get("protocol") or "").lower()
    if protocol not in PROTOCOLS:
        res.errors.append(f"{where}: protocol must be one of {', '.join(PROTOCOLS)}")
        return res

    data: dict[str, Any] = {"device_id": device_id}
    data["name"] = required_text(rec, "name", where, res)
    data["site"] = required_text(rec, "site", where, res)
    data["area"] = required_text(rec, "area", where, res)
    data["protocol"] = protocol
    enabled = as_bool(rec.get("enabled"), f"{where}.enabled", res)
    data["enabled"] = True if enabled is None else enabled
    put(data, "description", rec.get("description"))
    heartbeat = as_number(rec.get("history_heartbeat_s"), f"{where}.history_heartbeat_s", res,
                          integer=True)
    if heartbeat is not None:
        data["history_heartbeat_s"] = heartbeat

    poll = as_number(rec.get("poll_interval_ms"), f"{where}.poll_interval_ms", res, integer=True)
    if protocol == "mqtt":
        if poll is not None:
            res.errors.append(f"{where}: poll_interval_ms does not exist for mqtt devices")
    elif poll is not None:
        data["poll_interval_ms"] = poll

    conn: dict[str, Any] = {}
    placeholder_used = False

    def number(key: str, integer: bool = True) -> Any:
        return as_number(rec.get(key), f"{where}.{key}", res, integer=integer)

    if protocol == "modbus_tcp":
        if rec.get("host") in (None, TODO):
            conn["host"] = PLACEHOLDER_HOST
            placeholder_used = True
            res.todos.append(f"{where}: host unknown -> placeholder {PLACEHOLDER_HOST}")
        else:
            conn["host"] = str(rec["host"])
        put(conn, "port", number("port"))
        unit_id = number("unit_id")
        if unit_id is None:
            conn["unit_id"] = PLACEHOLDER_UNIT_ID
            placeholder_used = True
            res.todos.append(f"{where}: unit_id unknown -> placeholder {PLACEHOLDER_UNIT_ID}")
        else:
            conn["unit_id"] = unit_id
        if rec.get("framer") not in (None, TODO):
            conn["framer"] = str(rec["framer"]).lower()
        for key in ("timeout_ms", "retries"):
            put(conn, key, number(key))
        for key in ("byte_order", "word_order"):
            if rec.get(key) not in (None, TODO):
                conn[key] = str(rec[key]).lower()
        put(conn, "max_registers_per_read", number("max_registers_per_read"))
        group = tags_by_device["ModbusRegisters"].get(device_id, [])
        tags = modbus_tags(group, res)
    elif protocol == "modbus_rtu":
        conn["port"] = required_text(rec, "serial_port", where, res)
        unit_id = number("unit_id")
        if unit_id is None:
            res.errors.append(f"{where}: unit_id is required for modbus_rtu (1-247)")
        else:
            conn["unit_id"] = unit_id
        for key in ("baudrate", "bytesize", "stopbits", "timeout_ms", "retries",
                    "max_registers_per_read"):
            put(conn, key, number(key))
        if rec.get("parity") not in (None, TODO):
            conn["parity"] = str(rec["parity"]).upper()[:1]
        for key in ("byte_order", "word_order"):
            if rec.get(key) not in (None, TODO):
                conn[key] = str(rec[key]).lower()
        group = tags_by_device["ModbusRegisters"].get(device_id, [])
        tags = modbus_tags(group, res)
    elif protocol == "mqtt":
        conn["broker_ref"] = required_text(rec, "broker_ref", where, res)
        conn["telemetry_topic"] = required_text(rec, "telemetry_topic", where, res)
        for key in ("command_topic", "status_topic"):
            put(conn, key, rec.get(key))
        put(conn, "qos", number("qos"))
        put(conn, "timestamp_path", rec.get("timestamp_path"))
        if rec.get("timestamp_format") not in (None, TODO):
            conn["timestamp_format"] = str(rec["timestamp_format"]).lower()
        put(conn, "stale_after_s", number("stale_after_s", integer=False))
        tags = simple_tags(tags_by_device["MqttTags"].get(device_id, []), "MqttTags",
                           "json_path", res)
    elif protocol == "s7":
        if rec.get("host") in (None, TODO):
            conn["host"] = PLACEHOLDER_HOST
            placeholder_used = True
            res.todos.append(f"{where}: host unknown -> placeholder {PLACEHOLDER_HOST}")
        else:
            conn["host"] = str(rec["host"])
        for key in ("port", "rack", "slot", "timeout_ms"):
            put(conn, key, number(key))
        if rec.get("rack") in (None, TODO) or rec.get("slot") in (None, TODO):
            res.todos.append(f"{where}: rack/slot not given -> 0/1 (S7-1200/1500); S7-300 "
                             "uses 0/2")
        tags = simple_tags(tags_by_device["S7Tags"].get(device_id, []), "S7Tags", "address",
                           res, optional=("string_length",))
    elif protocol == "http":
        conn["base_url"] = required_text(rec, "base_url", where, res)
        auth = str(rec.get("auth_type") or "none").lower()
        if auth == "basic":
            conn["auth"] = {"type": "basic",
                            "username_env": required_text(rec, "username_env", where, res),
                            "password_env": required_text(rec, "password_env", where, res)}
        elif auth == "bearer":
            conn["auth"] = {"type": "bearer",
                            "token_env": required_text(rec, "token_env", where, res)}
        elif auth == "header":
            conn["auth"] = {"type": "header",
                            "header": required_text(rec, "auth_header", where, res),
                            "value_env": required_text(rec, "value_env", where, res)}
        elif auth != "none":
            res.errors.append(f"{where}: auth_type must be none, basic, bearer or header")
        verify = as_bool(rec.get("verify_tls"), f"{where}.verify_tls", res)
        if verify is False:
            conn["verify_tls"] = False
            res.todos.append(f"{where}: verify_tls false accepts any certificate")
        put(conn, "ca_file_env", rec.get("ca_file_env"))
        put(conn, "timeout_ms", number("timeout_ms"))
        put(conn, "timestamp_path", rec.get("timestamp_path"))
        if rec.get("timestamp_format") not in (None, TODO):
            conn["timestamp_format"] = str(rec["timestamp_format"]).lower()
        tags = simple_tags(tags_by_device["HttpTags"].get(device_id, []), "HttpTags",
                           "json_path", res, optional=("path",))
    elif protocol == "snmp":
        if rec.get("host") in (None, TODO):
            conn["host"] = PLACEHOLDER_HOST
            placeholder_used = True
            res.todos.append(f"{where}: host unknown -> placeholder {PLACEHOLDER_HOST}")
        else:
            conn["host"] = str(rec["host"])
        put(conn, "port", number("port"))
        version = str(rec.get("snmp_version") or "").lower().removeprefix("v")
        if version == "2c":
            conn["security"] = {"version": "2c",
                                "community_env": required_text(rec, "community_env", where,
                                                               res)}
        elif version == "3":
            sec: dict[str, Any] = {"version": "3",
                                   "username": required_text(rec, "snmp_user", where, res)}
            for key in ("auth_protocol", "priv_protocol"):
                if rec.get(key) not in (None, TODO):
                    sec[key] = str(rec[key]).lower()
            for key in ("auth_key_env", "priv_key_env"):
                put(sec, key, rec.get(key))
            conn["security"] = sec
        else:
            res.errors.append(f"{where}: snmp_version must be 2c or 3")
        for key in ("timeout_ms", "retries"):
            put(conn, key, number(key))
        tags = simple_tags(tags_by_device["SnmpTags"].get(device_id, []), "SnmpTags", "oid",
                           res)
    else:  # opcua
        if rec.get("endpoint_url") in (None, TODO):
            conn["endpoint_url"] = PLACEHOLDER_ENDPOINT
            placeholder_used = True
            res.todos.append(f"{where}: endpoint_url unknown -> placeholder")
        else:
            conn["endpoint_url"] = str(rec["endpoint_url"])
        conn["security_policy"] = required_text(rec, "security_policy", where, res)
        conn["security_mode"] = required_text(rec, "security_mode", where, res)
        auth = str(rec.get("auth_type") or "anonymous").lower()
        if auth == "username":
            conn["auth"] = {
                "type": "username",
                "username_env": required_text(rec, "username_env", where, res),
                "password_env": required_text(rec, "password_env", where, res),
            }
        elif auth == "anonymous":
            conn["auth"] = {"type": "anonymous"}
        else:
            res.errors.append(f"{where}: auth_type must be anonymous or username")
        for key in ("namespace", "client_certificate_env", "client_private_key_env"):
            put(conn, key, rec.get(key))
        put(conn, "timeout_ms", number("timeout_ms"))
        tags = simple_tags(tags_by_device["OpcuaTags"].get(device_id, []), "OpcuaTags",
                           "node_id", res)

    if placeholder_used:
        data["enabled"] = False
        res.todos.append(f"{where}: forced enabled: false until the placeholders are replaced")
    data["connection"] = conn
    data["tags"] = tags
    # Stable key order: identity, protocol-level settings, connection, tags.
    order = ["device_id", "name", "site", "area", "protocol", "enabled", "description",
             "poll_interval_ms", "history_heartbeat_s", "connection", "tags"]
    res.data = {k: data[k] for k in order if k in data}
    return res


def group_by_device(rows: list[tuple[int, dict[str, Any]]]) -> dict[str, list]:
    grouped: dict[str, list] = {}
    for row, rec in rows:
        did = rec.get("device_id")
        if did not in (None, TODO):
            grouped.setdefault(str(did).lower(), []).append((row, rec))
    return grouped


# ---------------------------------------------------------------------------------------------
# Output
# ---------------------------------------------------------------------------------------------


def render(res: Result, source: str) -> str:
    header = [f"# Generated from {source} by add-device/scripts/excel_to_devices.py.",
              "# Review before committing; run: python -m hub.validate_config"]
    if res.todos:
        header.append("# TODO:")
        header += [f"#   - {t}" for t in res.todos]
    body = yaml.safe_dump(res.data, sort_keys=False, allow_unicode=True, default_flow_style=False)
    return "\n".join(header) + "\n" + body


def prevalidate(candidates: list[Result], out_dir: Path, source: str) -> dict[str, list[str]]:
    """Validate candidates + existing config in a scratch copy; return {device_id: errors}.

    Uses the hub's own loader so the rules can never drift from ``hub.validate_config``.
    """
    if not candidates:
        return {}
    sys.path.insert(0, os.getcwd())
    try:
        from hub.config_loader import DEVICES_DIR, Severity, load_config
    except ImportError as exc:
        print(f"WARN    pre-write validation skipped (cannot import hub: {exc}); run from the "
              "project root with the project's Python >= 3.12, and use --check")
        return {}
    config_dir = out_dir.parent
    with tempfile.TemporaryDirectory() as tmp:
        scratch = Path(tmp) / "config"
        if config_dir.is_dir():
            shutil.copytree(config_dir, scratch)
        devices = scratch / DEVICES_DIR  # the loader only reads <config>/devices
        devices.mkdir(parents=True, exist_ok=True)
        new_sources = {}
        for res in candidates:
            (devices / f"{res.device_id}.yaml").write_text(render(res, source), encoding="utf-8")
            new_sources[f"{DEVICES_DIR}/{res.device_id}.yaml"] = res.device_id
        report = load_config(scratch, env={})
    rejected: dict[str, list[str]] = {}
    for issue in report.issues:
        if issue.severity is Severity.ERROR and issue.source in new_sources:
            where = f"{issue.location}: " if issue.location else ""
            rejected.setdefault(new_sources[issue.source], []).append(where + issue.message)
    return rejected


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--out", type=Path, default=Path("config/devices"),
                        help="output directory (default: %(default)s)")
    parser.add_argument("--only", nargs="+", metavar="DEVICE_ID", help="convert only these ids")
    parser.add_argument("--force", action="store_true", help="overwrite existing files")
    parser.add_argument("--dry-run", action="store_true", help="report only, write nothing")
    parser.add_argument("--check", action="store_true",
                        help="run python -m hub.validate_config on the config dir afterwards")
    args = parser.parse_args(argv)

    if not args.workbook.is_file():
        print(f"workbook not found: {args.workbook}", file=sys.stderr)
        return 2
    wb = openpyxl.load_workbook(args.workbook, data_only=True)
    if "Devices" not in wb.sheetnames:
        print("sheet 'Devices' not found", file=sys.stderr)
        return 2

    tags_by_device = {sheet: group_by_device(read_sheet(wb, sheet)) for sheet in TAG_SHEETS}
    device_rows = read_sheet(wb, "Devices")
    known_ids = {str(r.get("device_id")).lower() for _, r in device_rows if r.get("device_id")}
    failed = False
    for sheet, groups in tags_by_device.items():
        for did in groups.keys() - known_ids:
            print(f"ERROR   {sheet}: rows for unknown device_id {did!r} ignored")
            failed = True

    only = {d.lower() for d in args.only} if args.only else None
    seen: set[str] = set()
    candidates: list[Result] = []
    for row, rec in device_rows:
        res = build_device(row, rec, tags_by_device)
        if only and res.device_id not in only:
            continue
        if res.device_id in seen:
            res.errors.append(f"duplicate device_id {res.device_id!r} in workbook")
        seen.add(res.device_id)

        target = args.out / f"{res.device_id}.yaml"
        if not res.errors and target.exists() and not args.force:
            res.errors.append(f"{target} already exists (use --force to overwrite)")
        if res.errors or res.data is None:
            failed = True
            print(f"ERROR   {res.device_id}: not generated")
            for e in res.errors:
                print(f"          {e}")
            continue
        candidates.append(res)

    # Validate the candidates together with the existing config before touching any file.
    rejected = prevalidate(candidates, args.out, args.workbook.name)
    written = 0
    for res in candidates:
        target = args.out / f"{res.device_id}.yaml"
        if res.device_id in rejected:
            failed = True
            print(f"ERROR   {res.device_id}: not written - it would break the config:")
            for e in rejected[res.device_id]:
                print(f"          {e}")
            continue
        for e in res.todos:
            label = "WARN   " if "WRITABLE" in e else "TODO   "
            print(f"{label} {res.device_id}: {e}")
        n_tags = len(res.data.get("tags", []))
        if args.dry_run:
            print(f"DRY-RUN {target} ({n_tags} tags)")
        else:
            args.out.mkdir(parents=True, exist_ok=True)
            target.write_text(render(res, args.workbook.name), encoding="utf-8")
            print(f"WROTE   {target} ({n_tags} tags)")
            written += 1

    print(f"\n{written} file(s) written.")
    if args.check and not args.dry_run:
        config_dir = args.out.parent
        proc = subprocess.run(
            [sys.executable, "-m", "hub.validate_config", "--config-dir", str(config_dir)],
            check=False,
        )
        failed = failed or proc.returncode != 0
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
